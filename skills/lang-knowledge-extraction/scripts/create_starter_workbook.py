#!/usr/bin/env python3
import argparse
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation


CONCEPT_HEADERS = [
    "概念ID",
    "概念名称",
    "一句话定义",
    "边界/区别",
    "别名",
    "支撑判断ID",
    "来源",
    "状态",
    "可见范围",
    "审核备注",
    "创建时间",
    "更新时间",
    "系统写入键",
]

JUDGMENT_HEADERS = [
    "判断ID",
    "判断",
    "回应问题",
    "判断依据",
    "边界/改判条件",
    "关联概念ID",
    "来源",
    "状态",
    "可见范围",
    "审核备注",
    "创建时间",
    "更新时间",
    "系统写入键",
]


def style_sheet(ws, headers):
    ws.append(headers)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{ws.cell(1, len(headers)).column_letter}1000"
    ws.row_dimensions[1].height = 28

    fill = PatternFill("solid", fgColor="17324D")
    for cell in ws[1]:
        cell.fill = fill
        cell.font = Font(name="Sarasa UI SC", color="FFFFFF", bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")

    widths = [15, 24, 34, 34, 20, 22, 34, 16, 16, 28, 20, 20, 32]
    for index, width in enumerate(widths, start=1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width

    status = DataValidation(
        type="list",
        formula1='"pending_review,confirmed,discarded"',
        allow_blank=True,
    )
    access = DataValidation(
        type="list",
        formula1='"private,team,public_candidate,needs_review"',
        allow_blank=True,
    )
    ws.add_data_validation(status)
    ws.add_data_validation(access)
    status.add("H2:H1000")
    access.add("I2:I1000")

    pending_fill = PatternFill("solid", fgColor="FFF4CC")
    confirmed_fill = PatternFill("solid", fgColor="DDF3E4")
    discarded_fill = PatternFill("solid", fgColor="F1F1F1")
    ws.conditional_formatting.add("A2:M1000", FormulaRule(formula=['$H2="pending_review"'], fill=pending_fill))
    ws.conditional_formatting.add("A2:M1000", FormulaRule(formula=['$H2="confirmed"'], fill=confirmed_fill))
    ws.conditional_formatting.add("A2:M1000", FormulaRule(formula=['$H2="discarded"'], fill=discarded_fill))


def create_workbook(output: Path):
    wb = Workbook()
    guide = wb.active
    guide.title = "使用说明"
    guide.append(["知识萃取 starter 工作簿"])
    guide.append(["默认只使用概念库和判断库；候选先审核，再确认。"])
    guide.append(["状态", "pending_review＝待审核；confirmed＝已确认；discarded＝已废弃"])
    guide.append(["可见范围", "private＝私有；team＝团队；public_candidate＝公开候选；needs_review＝待判断"])
    guide.append(["来源", "填写原文件、链接、截图或录音位置；不要只留无出处摘要。"])
    guide.append(["关联", "支撑判断ID／关联概念ID填写稳定ID，不填写行号。"])
    guide.append(["配置", "把本文件的绝对路径作为 authority_ref；概念库与判断库分别映射到同名工作表。"])
    guide.column_dimensions["A"].width = 18
    guide.column_dimensions["B"].width = 86
    for row in guide.iter_rows():
        for cell in row:
            cell.font = Font(name="Sarasa UI SC")
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    guide["A1"].font = Font(name="Sarasa UI SC", size=16, bold=True, color="FFFFFF")
    guide["A1"].fill = PatternFill("solid", fgColor="17324D")
    guide.merge_cells("A1:B1")
    guide.freeze_panes = "A2"

    concepts = wb.create_sheet("概念库")
    judgments = wb.create_sheet("判断库")
    style_sheet(concepts, CONCEPT_HEADERS)
    style_sheet(judgments, JUDGMENT_HEADERS)

    wb.active = 0
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument("output", type=Path)
    args = parser.parse_args()
    create_workbook(args.output.resolve())


if __name__ == "__main__":
    main()
